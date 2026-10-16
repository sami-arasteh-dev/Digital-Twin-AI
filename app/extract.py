from __future__ import annotations
from pathlib import Path
import base64, mimetypes, io

def extract_file(path: str, mime: str):
    p=Path(path)
    suffix=p.suffix.lower()
    text=""
    if suffix in {".txt",".md",".csv",".json",".log",".xml",".html",".htm"}:
        text=p.read_text(encoding="utf-8",errors="ignore")
    elif suffix==".docx":
        try:
            from docx import Document
            doc=Document(str(p))
            text="\n".join(x.text for x in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        except Exception as e:
            text=f"[DOCX extraction error: {e}]"
    elif suffix==".pdf":
        try:
            from pypdf import PdfReader
            reader=PdfReader(str(p))
            text="\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:
            text=f"[PDF extraction error: {e}]"
    elif mime.startswith("image/"):
        try:
            from PIL import Image
            img=Image.open(str(p))
            text=f"[Image: {img.width}x{img.height}, mode={img.mode}]"
            try:
                import pytesseract
                ocr=pytesseract.image_to_string(img)
                if ocr.strip(): text += "\n[OCR]\n"+ocr
            except Exception:
                text += "\n[OCR unavailable]"
        except Exception as e:
            text=f"[Image processing error: {e}]"
    else:
        text=f"[Unsupported file type: {mime}]"
    return text

def image_data_url(path: str, mime: str):
    try:
        raw=Path(path).read_bytes()
        return f"data:{mime};base64,{base64.b64encode(raw).decode()}"
    except Exception:
        return None
